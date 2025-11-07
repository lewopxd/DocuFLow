#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Project: DocuFlow
File:  data_sheet_parser.py
Created: 2025-11-05
Author: @lewopxd

Description:
Parses data sheet files (xlsx) and extracts their
structure (metadata, sheets, tables, columns) or specific row data
into a JSON-serializable dictionary.
"""

import os
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple
import openpyxl
import openpyxl.worksheet.table
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils.cell import range_boundaries

# ============================================================
# CONFIGURACIÓN GLOBAL
# ============================================================
clean_row_text = True  # Si es True, elimina espacios al inicio y final de los valores de celda

# Importar el guardián de seguridad
try:
    # Importación relativa (sube a 'core', baja a 'file_helpers')
    from ..file_helpers.path_utils import create_safe_temp_copy
except ImportError as e:
    print(f"❌ CRITICAL: No se pudo importar 'create_safe_temp_copy' (ImportError: {e}).")
    from contextlib import contextmanager
    @contextmanager
    def create_safe_temp_copy(path):
        print("⚠️ USANDO FALLBACK NO SEGURO DE create_safe_temp_copy")
        if os.path.exists(path): yield path
        else: yield None

# -------------------------------------------------------------
# -------------------[   HELPER FUNCTIONS   ]------------------
# -------------------------------------------------------------

def _clean_cell_value(value: Any) -> Any:
    """
    Limpia el valor de una celda eliminando espacios al inicio y final
    si clean_row_text está activado y el valor es un string.
    
    Args:
        value: El valor de la celda (puede ser str, int, float, None, etc.)
        
    Returns:
        El valor limpio (sin espacios) si es string y clean_row_text=True,
        o el valor original en caso contrario.
    """
    global clean_row_text
    
    if clean_row_text and isinstance(value, str):
        return value.strip()
    return value

# -------------------------------------------------------------
# -------------------[   STRUCTURE PARSING   ]-----------------
# -------------------------------------------------------------

def get_data_sheet_structure(file_path: str) -> Optional[Dict[str, Any]]:
    """
    Función principal de entrada (Estructura).
    
    Analiza un archivo de hoja de datos (Excel) de forma segura
    y devuelve su estructura de metadatos.

    Args:
        file_path: La ruta completa al archivo original del usuario.

    Returns:
        Un diccionario con la estructura del archivo, o None si falla.
    """
    
    print(f"Solicitud de análisis de ESTRUCTURA para: {file_path}")
    
    with create_safe_temp_copy(file_path) as temp_path:
        
        if not temp_path:
            return None
            
        file_ext = Path(file_path).suffix.lower()
        original_filename = file_path 

        try:
            if file_ext == '.xlsx':
                return _parse_xlsx_structure(temp_path, original_filename)
                
            elif file_ext == '.xls':
                print(f"ℹ️ Parser para '.xls' no implementado aún.")
                return None
                
            elif file_ext == '.csv':
                print(f"ℹ️ Parser para '.csv' no implementado aún.")
                return None
                
            else:
                print(f"❌ Error: Tipo de archivo no soportado: {file_ext}")
                return None
                
        except Exception as e:
            print(f"❌ Error fatal durante el despacho de parsing de estructura: {e}")
            return None

def _parse_xlsx_structure(temp_file_path: str, original_filename: str) -> Optional[Dict[str, Any]]:
    """Parsea la ESTRUCTURA (metadatos) de un .xlsx."""
    workbook = None 
    try:
        workbook = openpyxl.load_workbook(temp_file_path, data_only=True)

        file_structure = {
            "metadata": _get_file_metadata(temp_file_path, original_filename, len(workbook.sheetnames)),
            "sheets": []
        }

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            sheet_data = {
                "sheetName": sheet_name,
                "tables": []
            }

            if sheet.tables:
                for table in sheet.tables.values():
                    table_data = _extract_table_structure_data(table, sheet)
                    sheet_data["tables"].append(table_data)
            
            file_structure["sheets"].append(sheet_data)

        workbook.close()
        return file_structure

    except InvalidFileException:
        print(f"❌ Error: El archivo '{original_filename}' no es un .xlsx válido.")
        return None
    except Exception as e:
        print(f"❌ Error inesperado parseando ESTRUCTURA .xlsx: {e}")
        import traceback
        traceback.print_exc()
        return None
    finally:
        if workbook:
            workbook.close()

def _extract_table_structure_data(table: openpyxl.worksheet.table.Table, 
                                  sheet: openpyxl.worksheet.worksheet.Worksheet) -> Dict[str, Any]:
    """Helper para extraer metadatos de una Tabla de openpyxl."""
    
    column_names = []
    
    try:
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        num_cols = (max_col - min_col) + 1
    except Exception:
        num_cols = 1 
        min_row, max_row = 1, 1 # Fallback

    if table.headerRowCount > 0:
        try:
            header_cells = sheet.iter_rows(
                min_row=min_row, 
                max_row=min_row + table.headerRowCount - 1, 
                min_col=min_col, 
                max_col=max_col,
                values_only=True
            )
            
            first_header_row = next(header_cells, [])
            column_names = [str(cell) if cell is not None else f"Columna_{i+1}" for i, cell in enumerate(first_header_row)]

        except Exception as e:
            print(f"⚠️ No se pudo parsear el encabezado de la tabla '{table.name}': {e}")
            column_names = [f"Columna_{i+1}" for i in range(num_cols)]
    else:
        column_names = [f"Columna_{i+1}" for i in range(num_cols)]
        
    data_row_count = (max_row - min_row + 1) - table.headerRowCount
    
    return {
        "tableName": table.name,
        "isNamedTable": True,
        "range": table.ref,
        "rowCount": data_row_count, # Devuelve solo las filas de DATOS
        "columns": column_names
    }

def _get_file_metadata(file_path: str, original_filename: str, sheet_count: int) -> Dict[str, Any]:
    """Construye el diccionario de metadatos."""
    try:
        size_bytes = os.path.getsize(file_path)
        size_kb = round(size_bytes / 1024, 2)
    except OSError:
        size_kb = 0

    return {
        "fullPath": original_filename, 
        "fileName": Path(original_filename).name,
        "fileSizeKB": size_kb,
        "sheetCount": sheet_count
    }

# --------------------------------------> END [ STRUCTURE PARSING ... ]


# -------------------------------------------------------------
# -------------------[   ROW DATA PARSING   ]------------------
# -------------------------------------------------------------

def _parse_range_config(range_config: dict, max_rows_in_table: int) -> Tuple[int, int]:
    """
    Parsea de forma segura el 'rangeConfig' y lo ajusta a los límites de la tabla.
    El rango es 1-based y se aplica a la tabla completa (fila 1 = encabezado).
    
    Args:
        range_config: El dict {"type": "...", "data": "..."}
        max_rows_in_table: El número total de filas en la tabla (incluyendo encabezado).
        
    Returns:
        Una tupla (start_row, end_row) de índices 1-based.
    """
    try:
        config_type = range_config.get("type", "all").lower()  # Default a "all"
        config_data = range_config.get("data")
        
        start_row = 1
        end_row = max_rows_in_table
        
        # === NUEVO: Tipo "all" ===
        if config_type == "all":
            # Desde la fila 2 (después del encabezado) hasta el final
            start_row = 2
            end_row = max_rows_in_table
            print(f"ℹ️ RangeConfig tipo 'all': Procesando desde fila 2 hasta {max_rows_in_table}")
            
        elif config_type == "unique_row":
            if config_data is None:
                print(f"⚠️ Error: tipo 'unique_row' requiere 'data'. Usando rango completo.")
                start_row = 2
                end_row = max_rows_in_table
            else:
                start_row = int(config_data)
                end_row = start_row
                
        elif config_type == "multiple_row":
            if config_data is None:
                print(f"⚠️ Error: tipo 'multiple_row' requiere 'data'. Usando rango completo.")
                start_row = 2
                end_row = max_rows_in_table
            else:
                parts = str(config_data).split('-')
                if len(parts) != 2:
                    print(f"⚠️ Error: formato de 'data' inválido para 'multiple_row': '{config_data}'. Usando rango completo.")
                    start_row = 2
                    end_row = max_rows_in_table
                else:
                    start_row = int(parts[0])
                    end_row = int(parts[1])
        else:
            print(f"⚠️ Tipo de rangeConfig desconocido: '{config_type}'. Usando rango completo ('all').")
            start_row = 2
            end_row = max_rows_in_table
            
    except (ValueError, TypeError, IndexError) as e:
        print(f"⚠️ Error parseando rangeConfig '{range_config}': {e}. Usando rango completo.")
        start_row = 2
        end_row = max_rows_in_table

    # Validar y ajustar límites
    if start_row < 1:
        start_row = 1
    if end_row > max_rows_in_table:
        end_row = max_rows_in_table
    if start_row > end_row:
        start_row, end_row = end_row, start_row # Invertir si están mal ordenados
        
    return start_row, end_row

def get_row_data_sheet(config: Dict[str, Any]) -> Optional[List[Dict[str, Any]]]:
    """
    Función principal de entrada (Datos de Filas).
    
    Extrae datos de filas específicas de una tabla y columnas específicas.
    El rango es 1-based (Fila 1 es el encabezado).

    Args:
        config: Un dict que DEBE contener:
            - "filePath": Ruta al archivo .xlsx
            - "sheetName": Nombre de la hoja
            - "tableName": Nombre de la tabla
            - "rangeConfig": {"type": "all"} o {"type": "multiple_row", "data": "1-100"} o {"type": "unique_row", "data": "5"}
                             Si no se especifica, por defecto es {"type": "all"}
            - "columns": ["Columna 1", "Columna 2"]

    Returns:
        Una lista de diccionarios (JSON) con los datos, o None si falla.
    """
    
    print(f"Solicitud de análisis de DATOS para: {config.get('filePath')}")
    
    # 1. Validar Configuración
    file_path = config.get("filePath")
    sheet_name = config.get("sheetName")
    table_name = config.get("tableName")
    
    # === NUEVO: rangeConfig es opcional, por defecto "all" ===
    range_config = config.get("rangeConfig", {"type": "all"})
    
    target_columns = config.get("columns")

    if not all([file_path, sheet_name, table_name, target_columns]):
        print(f"❌ Error: Configuración incompleta para get_row_data_sheet.")
        return None

    results_list = []
    workbook = None
    
    # 2. Copia Segura
    with create_safe_temp_copy(file_path) as temp_path:
        if not temp_path:
            return None
        
        try:
            # 3. Abrir Workbook
            workbook = openpyxl.load_workbook(temp_path, data_only=True)
            
            # 4. Localizar Hoja y Tabla
            if sheet_name not in workbook.sheetnames:
                print(f"❌ Error: Hoja '{sheet_name}' no encontrada.")
                workbook.close()
                return None
            sheet = workbook[sheet_name]
            
            if table_name not in sheet.tables:
                print(f"❌ Error: Tabla '{table_name}' no encontrada en la hoja '{sheet_name}'.")
                workbook.close()
                return None
            table = sheet.tables[table_name]
            
            # 5. Mapear Columnas y Límites de Tabla
            min_col, min_row_abs, max_col, max_row_abs = range_boundaries(table.ref)
            total_table_rows = (max_row_abs - min_row_abs) + 1
            
            # Obtener los nombres de encabezado reales (siempre en la primera fila de la tabla)
            header_cells = next(sheet.iter_rows(
                min_row=min_row_abs, 
                max_row=min_row_abs, 
                min_col=min_col, 
                max_col=max_col, 
                values_only=True
            ), [])
            actual_headers = [str(h) if h is not None else "" for h in header_cells]
            
            # Crear mapa de: {NombreDeColumnaSolicitada: indice_relativo_en_tabla (0-based)}
            column_map = {} # ej: {"Nombre": 1, "Apellido": 2}
            for name in target_columns:
                try:
                    rel_index = actual_headers.index(name)
                    column_map[name] = rel_index
                except ValueError:
                    print(f"⚠️ Advertencia: Columna '{name}' solicitada pero no encontrada en la tabla '{table_name}'.")
                    
            if not column_map:
                print(f"❌ Error: Ninguna de las columnas solicitadas fue encontrada.")
                workbook.close()
                return None

            # 6. Parsear Rango (con soporte para "all")
            start_req_idx, end_req_idx = _parse_range_config(range_config, total_table_rows)
            
            # 7. Calcular rangos absolutos para la iteración
            abs_start_row_to_iter = min_row_abs + start_req_idx - 1
            abs_end_row_to_iter = min_row_abs + end_req_idx - 1

            # 8. Extraer Datos (CON LIMPIEZA DE TEXTO)
            print(f"✓ Extrayendo filas {start_req_idx}-{end_req_idx} (absolutas {abs_start_row_to_iter}-{abs_end_row_to_iter}) de la tabla '{table_name}'...")
            
            for row_num, row_cells in enumerate(
                sheet.iter_rows(
                    min_row=abs_start_row_to_iter,
                    max_row=abs_end_row_to_iter,
                    min_col=min_col,
                    max_col=max_col,
                    values_only=True
                ), 
                start=start_req_idx # El índice 1-based que el usuario solicitó
            ):
                row_data = {"row_index": row_num}
                
                for col_name, rel_idx in column_map.items():
                    # rel_idx es el índice 0-based de la celda en el tuple row_cells
                    if rel_idx < len(row_cells):
                        # *** APLICAR LIMPIEZA DE TEXTO ***
                        raw_value = row_cells[rel_idx]
                        row_data[col_name] = _clean_cell_value(raw_value)
                    else:
                        row_data[col_name] = None # Seguridad por si el mapa está desfasado
                
                results_list.append(row_data)

            workbook.close()
            return results_list

        except Exception as e:
            print(f"❌ Error inesperado extrayendo datos de fila: {e}")
            import traceback
            traceback.print_exc()
            if workbook:
                workbook.close()
            return None

# --------------------------------------> END [ ROW DATA PARSING ... ]