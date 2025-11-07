#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Project: DocuFlow
File:  replace_holders_text.py
Created: 2025-11-05
Author: @lewopxd

Description:
Provides functions to safely replace text placeholders in 
Microsoft Word (.docx) documents.
"""

import os
from pathlib import Path
from typing import Dict, Any, Optional, Set
import docx
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

# Importar el guardián de seguridad
try:
    from ..file_helpers.path_utils import create_safe_temp_copy
except ImportError as e:
    print(f"❌ CRITICAL: No se pudo importar 'create_safe_temp_copy' (ImportError: {e}).")
    from contextlib import contextmanager
    @contextmanager
    def create_safe_temp_copy(path):
        print("⚠️ USANDO FALLBACK NO SEGURO DE create_safe_temp_copy")
        if os.path.exists(path): yield path
        else: yield None

# -------------------------------------------------------------
# -------------------[   CORE HELPER LOGIC   ]-----------------
# -------------------------------------------------------------

def _clear_run_highlight(run):
    """Quita el resaltado (highlight) de un Run."""
    run.font.highlight_color = None

def _perform_run_replacement(
    paragraph: Paragraph, 
    tag: str, 
    text_value: str, 
    clear_highlight: bool
) -> bool:
    """
    Realiza el reemplazo de texto a nivel de 'Run' para un tag específico
    en un párrafo.
    """
    
    # 1. Búsqueda simple (si el tag está completo en un solo run)
    for run in paragraph.runs:
        if tag in run.text:
            run.text = run.text.replace(tag, text_value)
            if clear_highlight:
                _clear_run_highlight(run)
            return True # Encontrado y reemplazado

    # 2. Búsqueda compleja (si el tag está dividido en varios runs)
    run_indices = []
    combined_text = ""
    
    for i, run in enumerate(paragraph.runs):
        combined_text += run.text
        run_indices.append(i)
        
        if not tag.startswith(combined_text):
            combined_text = ""
            run_indices = []
            
            if tag.startswith(run.text):
                combined_text = run.text
                run_indices = [i]

        if combined_text == tag:
            first_run_idx = run_indices[0]
            
            paragraph.runs[first_run_idx].text = text_value
            if clear_highlight:
                _clear_run_highlight(paragraph.runs[first_run_idx])
            
            for j in range(1, len(run_indices)):
                run_idx_to_clear = run_indices[j]
                paragraph.runs[run_idx_to_clear].text = ""
                if clear_highlight:
                    _clear_run_highlight(paragraph.runs[run_idx_to_clear])
            
            return True 
            
    return False 

def _process_paragraphs(
    paragraphs: list[Paragraph], 
    replacements: Dict[str, str], 
    clear_highlight: bool
) -> Set[str]:
    """Itera y aplica reemplazos a una lista de objetos de párrafo."""
    
    found_tags = set()
    
    for p in paragraphs:
        if '{{' not in p.text and '${{' not in p.text: # Chequeo rápido optimizado
            continue
            
        for tag, text_value in replacements.items():
            if tag in p.text:
                if _perform_run_replacement(p, tag, text_value, clear_highlight):
                    found_tags.add(tag)
                    
    return found_tags

def _process_tables(
    tables: list, 
    replacements: Dict[str, str], 
    clear_highlight: bool
) -> Set[str]:
    """Itera y aplica reemplazos a todas las celdas de una lista de tablas."""
    
    found_tags = set()
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                tags_in_cell = _process_paragraphs(
                    cell.paragraphs, 
                    replacements, 
                    clear_highlight
                )
                found_tags.update(tags_in_cell)
    return found_tags

# --------------------------------------> END [ CORE HELPER LOGIC ... ]

# -------------------------------------------------------------
# -------------------[   MAIN FUNCTION   ]---------------------
# -------------------------------------------------------------

def generate_document_from_template(config: Dict[str, Any]) -> Dict[str, Any]:
    """
    Función principal de entrada (Generador de Documentos).
    
    Reemplaza placeholders en un archivo .docx y lo guarda en un nuevo
    destino, devolviendo un log detallado del proceso.

    Args:
        config: Diccionario de configuración.
            - "source_path" (str)
            - "target_directory" (str): Carpeta base de salida
            - "target_filename" (str): Nombre (o ruta relativa) del archivo
            - "replacements" (Dict[str, str])
            - "overwrite" (bool, opcional)
            - "clear_highlight" (bool, opcional)
            - "author" (str, opcional): Nuevo autor del documento
            - "last_modified_by" (str, opcional): Nuevo modificador

    Returns:
        Un diccionario (JSON) con el estado del proceso.
    """
    
    status_report = {
        "success": False,
        "source_path": config.get("source_path"),
        "target_path": None,
        "error": None,
        "details": {
            "source_file_found": None,
            "target_directory_created": None,
            "file_overwritten": None,
            "metadata_updated": False, 
            "placeholder_status": {
                "all_found": False,
                "total_requested": len(config.get("replacements", {})),
                "total_found": 0,
                "found": [],
                "not_found": list(config.get("replacements", {}).keys())
            }
        }
    }
    
    replacements = config.get("replacements", {})
    source_path = config.get("source_path")
    target_dir = config.get("target_directory")
    target_filename = config.get("target_filename")
    
    try:
        if not all([source_path, target_dir, target_filename, replacements]):
            raise ValueError("Configuración incompleta (faltan source_path, target_directory, target_filename o replacements).")

        # 3. Validar y crear directorio de destino
        
        # --- [ INICIO DEL PARCHE ] ---
        # Construir la ruta final
        target_path = Path(target_dir) / target_filename
        status_report["target_path"] = str(target_path)
        
        # Obtener el directorio *padre* del archivo final
        dir_path = target_path.parent
        # --- [ FIN DEL PARCHE ] ---
        
        if not dir_path.exists():
            print(f"✓ Creando directorio de destino: {dir_path}")
            os.makedirs(dir_path, exist_ok=True)
            status_report["details"]["target_directory_created"] = True
        else:
            status_report["details"]["target_directory_created"] = False

        # 4. Validar 'overwrite'
        overwrite = config.get("overwrite", False)
        if target_path.exists() and not overwrite:
            raise FileExistsError(f"El archivo de destino '{target_path}' ya existe y 'overwrite' es False.")
        
        status_report["details"]["file_overwritten"] = target_path.exists() and overwrite

        # 5. Usar el Guardián de Seguridad para abrir la plantilla
        with create_safe_temp_copy(source_path) as temp_path:
            if not temp_path:
                status_report["details"]["source_file_found"] = False
                raise FileNotFoundError("Error de validación: El archivo fuente no existe o no es un archivo.")
            
            status_report["details"]["source_file_found"] = True
            
            doc: DocxDocument = docx.Document(temp_path)
            clear_highlight = config.get("clear_highlight", False)
            all_found_tags = set()

            # 6. Iterar y reemplazar
            for section in doc.sections:
                all_found_tags.update(_process_paragraphs(
                    section.header.paragraphs, replacements, clear_highlight
                ))
                all_found_tags.update(_process_tables(
                    section.header.tables, replacements, clear_highlight
                ))
                
            all_found_tags.update(_process_paragraphs(
                doc.paragraphs, replacements, clear_highlight
            ))
            all_found_tags.update(_process_tables(
                doc.tables, replacements, clear_highlight
            ))

            for section in doc.sections:
                all_found_tags.update(_process_paragraphs(
                    section.footer.paragraphs, replacements, clear_highlight
                ))
                all_found_tags.update(_process_tables(
                    section.footer.tables, replacements, clear_highlight
                ))

            
            # --- Actualización de Metadatos ---
            new_author = config.get("author")
            new_last_modified = config.get("last_modified_by")
            
            if new_author:
                doc.core_properties.author = new_author
                status_report["details"]["metadata_updated"] = True
            
            if new_last_modified:
                doc.core_properties.last_modified_by = new_last_modified
                status_report["details"]["metadata_updated"] = True
            
            
            # 7. Guardar el nuevo documento
            doc.save(target_path)
            
            # 8. Actualizar el log de estado
            found_list = sorted(list(all_found_tags))
            not_found_list = sorted(list(set(replacements.keys()) - all_found_tags))
            
            status_report["success"] = True
            status_report["details"]["placeholder_status"]["total_found"] = len(found_list)
            status_report["details"]["placeholder_status"]["found"] = found_list
            status_report["details"]["placeholder_status"]["not_found"] = not_found_list
            status_report["details"]["placeholder_status"]["all_found"] = (len(not_found_list) == 0)

    except Exception as e:
        status_report["error"] = str(e)
        status_report["success"] = False
    
    return status_report

# --------------------------------------> END [ MAIN FUNCTION ... ]