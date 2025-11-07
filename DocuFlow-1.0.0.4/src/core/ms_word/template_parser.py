#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Project: DocuFlow
File:    template_parser.py
Created: 2025-11-05
Author:  @lewopxd

Description:
Provides functions to safely scan and parse Word (.docx) templates,
extracting metadata and a unique list of placeholders.
"""

import os
import re
from pathlib import Path
from typing import Dict, Any, Optional, List, Set, Tuple
import docx
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

# Importar el guardián de seguridad
try:
    from ..file_helpers.path_utils import create_safe_temp_copy
except ImportError as e:
    print(f"❌ CRITICAL: No se pudo importar 'create_safe_temp_copy' (ImportError: {e}).")
    from contextlib import contextmanager
    @contextmanager
    def create_safe_temp_copy(path):
        print("⚠️ USANDO FALLBACK NO SEGURO DE create_safe_temp_copy")
        if os.path.exists(path): yield path
        else: yield None

# -------------------------------------------------------------
# -------------------[   CORE HELPER LOGIC   ]-----------------
# -------------------------------------------------------------

def _build_regex_map(definitions: List[Dict[str, str]]) -> Dict[str, re.Pattern]:
    """
    Construye un mapa de Expresiones Regulares compiladas a partir de
    las definiciones de placeholders.
    """
    regex_map = {}
    for definition in definitions:
        try:
            prefix = re.escape(definition["prefix"])
            suffix = re.escape(definition["suffix"])
            
            # --- [ INICIO DE CORRECCIÓN ] ---
            # El bug estaba en la regex anterior: [a-zA-Z0-9_]+
            # Esa regex no permitía caracteres Unicode como 'Ñ' o 'Í'.
            #
            # La nueva regex: [^\s]+?
            # Significa:
            # [^\s]  -> "Cualquier carácter que NO sea un espacio en blanco"
            # +?     -> "Una o más veces, de forma no-codiciosa (non-greedy)"
            #
            # Esto permite {{AÑO_POSTULACION}} y {{DÍA_HV}} sin problemas.
            
            pattern_str = f"({prefix}[^\\s]+?{suffix})"
            # --- [ FIN DE CORRECCIÓN ] ---
            
            regex_map[definition["type"]] = re.compile(pattern_str)
            
        except (re.error, TypeError, KeyError) as e:
            print(f"⚠️ Error al construir Regex para la definición {definition}: {e}")
    return regex_map

def _scan_paragraphs(
    paragraphs: List[Paragraph], 
    regex_map: Dict[str, re.Pattern]
) -> Dict[str, Set[str]]:
    """
    Escanea una lista de párrafos usando la lógica 'Run-Stitcher'
    para encontrar todos los placeholders.
    """
    found_placeholders = {type_name: set() for type_name in regex_map}
    
    for p in paragraphs:
        # Usar p.text en lugar de join(run.text) es más simple y
        # robusto para la simple *detección* de texto.
        combined_text = p.text
        
        if not combined_text:
            continue

        for type_name, regex_pattern in regex_map.items():
            try:
                matches = regex_pattern.findall(combined_text)
                if matches:
                    found_placeholders[type_name].update(matches)
            except re.error as e:
                print(f"Error aplicando Regex en párrafo: {e}")
                
    return found_placeholders

def _extract_metadata(doc: DocxDocument, source_path: str) -> Dict[str, Any]:
    """Extrae metadatos básicos del documento."""
    
    # Convertir datetimes a strings ISO para que sean serializables en JSON
    # También añadimos chequeos por si la propiedad es None
    
    created_date = None
    if doc.core_properties.created:
        created_date = doc.core_properties.created.isoformat()
        
    modified_date = None
    if doc.core_properties.modified:
        modified_date = doc.core_properties.modified.isoformat()

    metadata = {
        "source_path": source_path,
        "file_size_kb": 0,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "author": doc.core_properties.author,
        "last_modified_by": doc.core_properties.last_modified_by,
        "created": created_date,
        "modified": modified_date,
        "title": doc.core_properties.title,
        "subject": doc.core_properties.subject,
    }
    
    try:
        size_bytes = os.path.getsize(source_path)
        metadata["file_size_kb"] = round(size_bytes / 1024, 2)
    except OSError:
        pass 
        
    return metadata

# --------------------------------------> END [ CORE HELPER LOGIC ... ]

# -------------------------------------------------------------
# -------------------[   MAIN FUNCTION   ]---------------------
# -------------------------------------------------------------

def get_template_info(config: Dict[str, Any]) -> Dict[str, Any]:
    """
    Función principal de entrada (Inspector de Plantillas).
    
    Analiza una plantilla .docx de forma segura, extrae metadatos
    y todos los placeholders únicos que coincidan con las definiciones.

    Args:
        config: Diccionario de configuración.
            - "source_path" (str): Ruta a la plantilla .docx
            - "definitions" (List[dict]): Lista de tipos de placeholders.
              Ej: [{"type": "text", "prefix": "${{", "suffix": "}}"}, ...]

    Returns:
        Un diccionario (JSON) con los metadatos y placeholders.
    """
    
    status_report = {
        "success": False,
        "error": None,
        "metadata": None,
        "placeholders": {},
        "stats": {
            "total_unique_tags": 0,
            "all_tags_found": []
        }
    }
    
    source_path = config.get("source_path")
    definitions = config.get("definitions")

    try:
        if not source_path or not definitions:
            raise ValueError("Configuración incompleta (faltan 'source_path' o 'definitions').")
            
        regex_map = _build_regex_map(definitions)
        if not regex_map:
            raise ValueError("No se pudieron construir patrones Regex a partir de 'definitions'.")

        all_found_tags = set()
        placeholders_by_type = {type_name: set() for type_name in regex_map}
        
        with create_safe_temp_copy(source_path) as temp_path:
            if not temp_path:
                raise FileNotFoundError("Error de validación: El archivo fuente no existe o no es un archivo.")
            
            doc: DocxDocument = docx.Document(temp_path)
            
            status_report["metadata"] = _extract_metadata(doc, source_path)

            # --- Encabezados ---
            for section in doc.sections:
                found_in_header = _scan_paragraphs(section.header.paragraphs, regex_map)
                for type_name, tags_set in found_in_header.items():
                    placeholders_by_type[type_name].update(tags_set)
                
                for table in section.header.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            found_in_cell = _scan_paragraphs(cell.paragraphs, regex_map)
                            for type_name, tags_set in found_in_cell.items():
                                placeholders_by_type[type_name].update(tags_set)

            # --- Cuerpo (Párrafos y Tablas) ---
            found_in_body = _scan_paragraphs(doc.paragraphs, regex_map)
            for type_name, tags_set in found_in_body.items():
                placeholders_by_type[type_name].update(tags_set)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        found_in_cell = _scan_paragraphs(cell.paragraphs, regex_map)
                        for type_name, tags_set in found_in_cell.items():
                            placeholders_by_type[type_name].update(tags_set)
            
            # --- Pies de página ---
            for section in doc.sections:
                found_in_footer = _scan_paragraphs(section.footer.paragraphs, regex_map)
                for type_name, tags_set in found_in_footer.items():
                    placeholders_by_type[type_name].update(tags_set)

                for table in section.footer.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            found_in_cell = _scan_paragraphs(cell.paragraphs, regex_map)
                            for type_name, tags_set in found_in_cell.items():
                                placeholders_by_type[type_name].update(tags_set)
            
            status_report["success"] = True
            
            for type_name, tags_set in placeholders_by_type.items():
                sorted_tags = sorted(list(tags_set))
                status_report["placeholders"][type_name] = sorted_tags
                all_found_tags.update(sorted_tags)
                
            status_report["stats"]["all_tags_found"] = sorted(list(all_found_tags))
            status_report["stats"]["total_unique_tags"] = len(all_found_tags)

    except Exception as e:
        status_report["error"] = str(e)
        status_report["success"] = False

    return status_report

# --------------------------------------> END [ MAIN FUNCTION ... ]